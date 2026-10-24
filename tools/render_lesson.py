# -*- coding: utf-8 -*-
"""Render one lesson to HTML + DOCX.

Usage:  python tools/render_lesson.py <lesson-dir>        e.g. 03_Strings

A lesson directory must contain:
  content.py    -- SLUG, TITLE, SUBTITLE, BRAND, FOOTER, SECTIONS,
                   CHEATSHEET, PITFALLS, COMPLEXITY
  problems.json -- {"group": "D", "groups": [{key,name,maps}], "problems": [...]}

Section numbers for the four trailing chapters (cheat sheet, pitfalls,
complexity, practice) are derived from len(SECTIONS), so a lesson may carry any
number of theory sections.
"""
import json, re, sys, html, pathlib, argparse, importlib.util

sys.stdout.reconfigure(encoding='utf-8')

ROOT = pathlib.Path(__file__).resolve().parent.parent

_ap = argparse.ArgumentParser()
_ap.add_argument("lesson", help="lesson directory, e.g. 03_Strings")
_ap.add_argument("--html-only", action="store_true", help="skip the DOCX build")
_args = _ap.parse_args()

OUT = (ROOT / _args.lesson).resolve()
if not (OUT / "content.py").exists():
    sys.exit("no content.py in %s" % OUT)

_spec = importlib.util.spec_from_file_location("lesson_content", OUT / "content.py")
C = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(C)

_pj    = json.loads((OUT / "problems.json").read_text(encoding="utf-8"))
SEL    = _pj["problems"]
GROUPS = [(g["key"], g["name"], g["maps"]) for g in _pj["groups"]]
PGROUP = _pj.get("group", "D")
SLUG   = getattr(C, "SLUG", OUT.name)
BRAND  = getattr(C, "BRAND", SLUG.replace("_", " "))
FOOTER = getattr(C, "FOOTER", "Problem links point to arena.olimpiici.com; statements are in "
                              "Bulgarian and downloadable as PDF from each problem page.")

PIT_T  = getattr(C, "PITFALLS_TITLE", "%d pitfalls that cost real marks" % len(C.PITFALLS))
NSEC   = len(C.SECTIONS)
N_CHE, N_PIT, N_CPX, N_PRB = NSEC + 1, NSEC + 2, NSEC + 3, NSEC + 4
NPROB  = len(SEL)
_syl   = json.loads((ROOT / "syllabus.json").read_text(encoding="utf-8"))
ORDER  = [l for t in _syl["tracks"] for l in t["lessons"]
          if (ROOT / l["slug"] / "content.py").exists()]
_pos   = next((i for i, l in enumerate(ORDER) if l["slug"] == SLUG), None)
PREV   = ORDER[_pos - 1] if _pos else None
NEXT   = ORDER[_pos + 1] if _pos is not None and _pos + 1 < len(ORDER) else None
ROUND_EN = {"Общински кръг": "Municipal round", "Областен кръг": "Regional round",
            "Национален кръг": "National round", "Есенен турнир": "Autumn tournament",
            "Зимен турнир": "Winter tournament", "Пролетен турнир": "Spring tournament",
            "Летен турнир": "Summer tournament"}

# --------------------------------------------------------------- inline markup
TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`)', re.S)

def inline(text):
    """-> list of (kind, text) with kind in {'t','c','b'}"""
    out = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            out.append(('b', part[2:-2]))
        elif part.startswith('`') and part.endswith('`'):
            out.append(('c', part[1:-1]))
        else:
            out.append(('t', part))
    return out

def h_inline(text):
    buf = []
    for kind, t in inline(text):
        e = html.escape(t)
        buf.append(e if kind == 't' else
                   ('<strong>%s</strong>' % e if kind == 'b' else '<code>%s</code>' % e))
    return ''.join(buf)

def stars(n):
    return '\u2605' * n + '\u2606' * (5 - n)

# =============================================================== HTML
CSS = """
:root{
  --bg:#fbfaf8; --panel:#ffffff; --ink:#1c1b1a; --muted:#6b6763; --line:#e4e0da;
  --accent:#8a4b1f; --accent-soft:#f3ece4; --code-bg:#f6f3ee; --code-ink:#2d2a26;
  --kw:#8a4b1f; --str:#3f6b3f; --com:#9a948c; --num:#2f5c8a; --pre:#7a5ea8;
  --note-bg:#fdf6e7; --note-line:#d9a441; --ok:#3f6b3f;
}
:root:not([data-theme="light"]){ }
@media (prefers-color-scheme: dark){
  :root:not([data-theme="light"]){
    --bg:#17161a; --panel:#1e1d22; --ink:#e9e6e1; --muted:#9a948c; --line:#33313a;
    --accent:#e0a06a; --accent-soft:#2a2530; --code-bg:#232128; --code-ink:#ddd8d1;
    --kw:#e0a06a; --str:#9ac48a; --com:#7b757f; --num:#84b4dd; --pre:#c0a2e8;
    --note-bg:#2b2618; --note-line:#a8813a; --ok:#9ac48a;
  }
}
:root[data-theme="dark"]{
  --bg:#17161a; --panel:#1e1d22; --ink:#e9e6e1; --muted:#9a948c; --line:#33313a;
  --accent:#e0a06a; --accent-soft:#2a2530; --code-bg:#232128; --code-ink:#ddd8d1;
  --kw:#e0a06a; --str:#9ac48a; --com:#7b757f; --num:#84b4dd; --pre:#c0a2e8;
  --note-bg:#2b2618; --note-line:#a8813a; --ok:#9ac48a;
}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);
  font:16px/1.65 "Iowan Old Style","Palatino Linotype",Palatino,Georgia,serif;}
code,pre,.mono{font-family:"JetBrains Mono",Consolas,"SF Mono",Menlo,monospace}
.wrap{display:flex;gap:0;align-items:flex-start;max-width:1220px;margin:0 auto}
/* ---- sidebar ---- */
#nav{position:sticky;top:0;flex:0 0 250px;height:100vh;overflow-y:auto;
  padding:26px 14px 40px 20px;border-right:1px solid var(--line);font-size:13.5px;
  font-family:system-ui,-apple-system,"Segoe UI",sans-serif}
#nav .brand{font-weight:700;letter-spacing:.02em;color:var(--accent);
  font-size:12px;text-transform:uppercase;margin-bottom:14px}
#nav a{display:block;padding:4px 8px;margin:1px 0;border-radius:5px;
  color:var(--muted);text-decoration:none;line-height:1.35}
#nav a:hover{background:var(--accent-soft);color:var(--ink)}
#nav a.on{background:var(--accent-soft);color:var(--accent);font-weight:600}
#nav .grp{margin-top:14px;font-size:11px;text-transform:uppercase;letter-spacing:.06em;
  color:var(--muted);opacity:.7;padding:0 8px}
#nav a.home{display:inline-flex;align-items:center;gap:7px;margin:0 0 16px;padding:5px 12px;
  border:1px solid var(--line);border-radius:6px;background:var(--panel);color:var(--ink);
  font-weight:600}
#nav a.home svg{width:15px;height:15px}
#nav a.home:hover{border-color:var(--accent);color:var(--accent);background:var(--panel)}
#nav .pager{display:flex;gap:8px;margin-top:22px;padding-top:16px;border-top:1px solid var(--line)}
#nav .pager a,#nav .pager span{flex:1 1 0;display:flex;flex-direction:column;align-items:center;
  margin:0;padding:6px 4px 7px;border:1px solid var(--line);border-radius:6px;
  background:var(--panel);color:var(--ink);font-size:22px;line-height:1.1}
#nav .pager small{font-size:11px;color:var(--muted);margin-top:2px}
#nav .pager a:hover{border-color:var(--accent);color:var(--accent);background:var(--panel)}
#nav .pager a:hover small{color:var(--accent)}
#nav .pager span{opacity:.35}
/* ---- main ---- */
main{flex:1 1 auto;min-width:0;padding:34px 40px 120px;max-width:900px}
header.top{border-bottom:2px solid var(--accent);padding-bottom:18px;margin-bottom:8px}
h1{font-size:30px;line-height:1.2;margin:0 0 6px}
.sub{color:var(--muted);font-size:14.5px;font-family:system-ui,sans-serif}
h2{font-size:21px;margin:44px 0 12px;padding-top:10px;border-top:1px solid var(--line);
  scroll-margin-top:16px}
h2:first-of-type{border-top:none}
h3{font-size:16.5px;margin:26px 0 8px;color:var(--accent);scroll-margin-top:16px}
p{margin:11px 0}
code{background:var(--code-bg);padding:1.5px 5px;border-radius:4px;font-size:.87em;
  color:var(--code-ink);border:1px solid var(--line)}
pre{background:var(--code-bg);border:1px solid var(--line);border-left:3px solid var(--accent);
  border-radius:7px;padding:13px 15px;overflow-x:auto;font-size:13px;line-height:1.55;
  margin:14px 0;color:var(--code-ink)}
pre code{background:none;border:none;padding:0;font-size:inherit;color:inherit}
.kw{color:var(--kw);font-weight:600}.str{color:var(--str)}.com{color:var(--com);font-style:italic}
.num{color:var(--num)}.pp{color:var(--pre)}
table{border-collapse:collapse;width:100%;margin:15px 0;font-size:14px;
  font-family:system-ui,-apple-system,"Segoe UI",sans-serif}
th,td{text-align:left;padding:7px 10px;border-bottom:1px solid var(--line);vertical-align:top}
th{font-size:11.5px;text-transform:uppercase;letter-spacing:.05em;color:var(--muted);
  border-bottom:1.5px solid var(--line);font-weight:600}
tbody tr:nth-child(even){background:var(--accent-soft)}
ul{margin:11px 0;padding-left:22px}
li{margin:4px 0}
.note{background:var(--note-bg);border-left:3px solid var(--note-line);border-radius:0 7px 7px 0;
  padding:11px 15px;margin:15px 0;font-size:14.5px}
.tblwrap{overflow-x:auto}
/* ---- problems ---- */
.pgroup{margin:26px 0 0}
.pgroup h3{margin-bottom:2px}
.pgroup .maps{color:var(--muted);font-size:12.5px;font-family:system-ui,sans-serif;margin:0 0 8px}
table.prob td{padding:9px 10px}
table.prob a{color:var(--accent);text-decoration:none;font-weight:600;
  font-family:Consolas,monospace;font-size:13.5px}
table.prob a:hover{text-decoration:underline}
.tech{color:var(--muted);font-size:13px}
.src{color:var(--muted);font-size:12px;white-space:nowrap}
.dif{color:var(--accent);letter-spacing:1px;font-size:11.5px;white-space:nowrap}
.done td{opacity:.45}
.done td a{text-decoration:line-through}
input[type=checkbox]{width:16px;height:16px;accent-color:var(--accent);cursor:pointer}
#bar{position:sticky;top:0;z-index:5;background:var(--bg);padding:10px 0 8px;
  border-bottom:1px solid var(--line);display:flex;gap:12px;align-items:center;
  font-family:system-ui,sans-serif;font-size:13px;flex-wrap:wrap}
#bar .track{flex:1 1 130px;height:7px;background:var(--line);border-radius:4px;overflow:hidden;
  min-width:110px}
#bar .fill{height:100%;background:var(--accent);width:0;transition:width .25s}
button{font:inherit;font-size:12.5px;padding:4px 11px;border:1px solid var(--line);
  border-radius:6px;background:var(--panel);color:var(--ink);cursor:pointer}
button:hover{border-color:var(--accent);color:var(--accent)}
#themeBtn{position:fixed;top:14px;right:16px;z-index:10}
.cheat td:first-child{font-family:system-ui,sans-serif;width:33%}
.pit b{display:block;margin-bottom:2px}
footer{margin-top:60px;padding-top:16px;border-top:1px solid var(--line);
  color:var(--muted);font-size:12.5px;font-family:system-ui,sans-serif}
@media(max-width:900px){
  #nav{display:none}
  main{padding:26px 18px 90px}
  h1{font-size:24px}
  #themeBtn{top:8px;right:10px}
}
"""

JS = r"""
var THEME_KEY="cpp_lessons_theme";
var DONE_KEY="cpp_done_"+LESSON_SLUG;
// ---- theme -------------------------------------------------------------
var tb=document.getElementById('themeBtn');
function setTheme(t){
  if(t) document.documentElement.setAttribute('data-theme',t);
  else document.documentElement.removeAttribute('data-theme');
  try{t?localStorage.setItem(THEME_KEY,t):localStorage.removeItem(THEME_KEY)}catch(e){}
}
try{var st=localStorage.getItem(THEME_KEY); if(st) setTheme(st);}catch(e){}
tb.onclick=function(){
  var cur=document.documentElement.getAttribute('data-theme');
  if(!cur){cur=matchMedia('(prefers-color-scheme: dark)').matches?'dark':'light';}
  setTheme(cur==='dark'?'light':'dark');
};

// ---- progress ----------------------------------------------------------
var boxes=[].slice.call(document.querySelectorAll('input[type=checkbox][data-pid]'));
var fill=document.querySelector('#bar .fill'), lbl=document.getElementById('cnt');
// state lives in memory; localStorage is a best-effort mirror so that ticking
// still works when storage is blocked (private window, file:// restrictions)
var state={};
try{
  var raw=localStorage.getItem(DONE_KEY);
  if(raw===null && LEGACY_KEY) raw=localStorage.getItem(LEGACY_KEY);  // one-time adopt
  state=JSON.parse(raw||'{}')||{};
}catch(e){state={}}
function save(){try{localStorage.setItem(DONE_KEY,JSON.stringify(state))}catch(e){}}
function paint(){
  var n=0;
  boxes.forEach(function(b){
    var on=!!state[b.dataset.pid]; b.checked=on; if(on)n++;
    b.closest('tr').classList.toggle('done',on);
  });
  fill.style.width=(n/boxes.length*100)+'%';
  lbl.textContent=n+' / '+boxes.length+' solved';
}
boxes.forEach(function(b){b.onchange=function(){
  if(b.checked)state[b.dataset.pid]=1; else delete state[b.dataset.pid];
  save(); paint();
};});
document.getElementById('reset').onclick=function(){
  if(confirm('Clear all progress?')){state={};save();paint();}
};
paint();

// ---- scroll spy --------------------------------------------------------
var links=[].slice.call(document.querySelectorAll('#nav a[href^="#"]'));
var map={}; links.forEach(function(a){
  var el=document.getElementById(a.getAttribute('href').slice(1)); if(el)map[el.id]=a;});
var obs=new IntersectionObserver(function(es){
  es.forEach(function(e){
    if(e.isIntersecting){
      links.forEach(function(a){a.classList.remove('on')});
      if(map[e.target.id])map[e.target.id].classList.add('on');
    }});
},{rootMargin:'0px 0px -75% 0px'});
Object.keys(map).forEach(function(id){obs.observe(document.getElementById(id))});

// ---- C++ highlighter ---------------------------------------------------
var KW=('alignas|alignof|auto|bool|break|case|catch|char|class|const|constexpr|continue|'+
 'decltype|default|delete|do|double|else|enum|explicit|export|extern|false|float|for|friend|'+
 'goto|if|inline|int|long|mutable|namespace|new|noexcept|nullptr|operator|private|protected|'+
 'public|register|return|short|signed|sizeof|static|static_cast|struct|switch|template|this|'+
 'throw|true|try|typedef|typename|union|unsigned|using|virtual|void|volatile|while|'+
 'string|string_view|vector|size_t|map|unordered_map|set|pair|istringstream|ostringstream').split('|');
var KWRE=new RegExp('\\b('+KW.join('|')+')\\b','g');
document.querySelectorAll('pre code').forEach(function(el){
  var src=el.textContent, out='', i=0;
  function esc(s){return s.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');}
  while(i<src.length){
    var rest=src.slice(i), m;
    if((m=rest.match(/^\/\/[^\n]*/))){out+='<span class="com">'+esc(m[0])+'</span>';i+=m[0].length;continue;}
    if((m=rest.match(/^#[a-z]+/))){out+='<span class="pp">'+esc(m[0])+'</span>';i+=m[0].length;continue;}
    if((m=rest.match(/^"(?:\\.|[^"\\])*"/))||(m=rest.match(/^'(?:\\.|[^'\\])*'/))){
      out+='<span class="str">'+esc(m[0])+'</span>';i+=m[0].length;continue;}
    if((m=rest.match(/^\b\d[\w.]*/))){out+='<span class="num">'+esc(m[0])+'</span>';i+=m[0].length;continue;}
    if((m=rest.match(/^[A-Za-z_]\w*/))){
      out+= KWRE.test(m[0])&&(KWRE.lastIndex=0,true) ? '<span class="kw">'+esc(m[0])+'</span>' : esc(m[0]);
      i+=m[0].length;continue;}
    out+=esc(src[i]); i++;
  }
  el.innerHTML=out;
});
"""


def build_html():
    P = []
    a = P.append
    a('<!DOCTYPE html>')
    a('<html lang="en"><head>')
    a('<meta charset="utf-8">')
    a('<meta name="viewport" content="width=device-width,initial-scale=1">')
    a('<title>%s</title>' % html.escape(C.TITLE))
    a('<style>%s</style>' % CSS)
    a('</head><body>')
    a('<button id="themeBtn" title="Toggle light / dark">&#9681;</button>')
    a('<div class="wrap">')

    # ---------------- sidebar
    a('<nav id="nav"><a class="home" href="../index.html" title="All lessons">'
      '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" '
      'stroke-linecap="round" stroke-linejoin="round"><path d="M3 11l9-8 9 8"/>'
      '<path d="M5 9.5V21h5v-6h4v6h5V9.5"/></svg>Home</a>')
    a('<div class="brand">%s</div>' % html.escape(BRAND))
    a('<div class="grp">Theory</div>')
    for sid, title, _ in C.SECTIONS:
        a('<a href="#%s">%s</a>' % (sid, html.escape(title)))
    a('<div class="grp">Reference</div>')
    a('<a href="#cheat">%d. Cheat sheet</a>' % N_CHE)
    a('<a href="#pitfalls">%d. Pitfalls</a>' % N_PIT)
    a('<a href="#complexity">%d. Complexity</a>' % N_CPX)
    a('<div class="grp">Practice</div>')
    a('<a href="#problems">%d. %d problems</a>' % (N_PRB, NPROB))
    for k, name, _ in GROUPS:
        a('<a href="#g%s" style="padding-left:16px">%s</a>' % (k, html.escape(name)))
    if _pos is not None:
        a('<div class="pager">')
        for l, arrow, word in ((PREV, '&larr;', 'Previous'), (NEXT, '&rarr;', 'Next')):
            if l:
                a('<a href="../%s/%s.html" title="%s: Lesson %s &middot; %s">%s<small>Lesson %s</small></a>'
                  % (l["slug"], l["slug"], word, l["num"], html.escape(l["title"]), arrow, l["num"]))
            else:
                a('<span>%s<small>&nbsp;</small></span>' % arrow)
        a('</div>')
    a('</nav>')

    # ---------------- main
    a('<main>')
    a('<header class="top"><h1>%s</h1><div class="sub">%s</div></header>'
      % (html.escape(C.TITLE), html.escape(C.SUBTITLE)))

    for sid, title, blocks in C.SECTIONS:
        a('<h2 id="%s">%s</h2>' % (sid, html.escape(title)))
        for kind, *rest in blocks:
            if kind == 'p':
                a('<p>%s</p>' % h_inline(rest[0]))
            elif kind == 'h3':
                a('<h3>%s</h3>' % h_inline(rest[0]))
            elif kind == 'code':
                a('<pre><code>%s</code></pre>' % html.escape(rest[0]))
            elif kind == 'note':
                a('<div class="note">%s</div>' % h_inline(rest[0]))
            elif kind == 'list':
                a('<ul>' + ''.join('<li>%s</li>' % h_inline(x) for x in rest[0]) + '</ul>')
            elif kind == 'table':
                head, rows = rest
                a('<div class="tblwrap"><table><thead><tr>')
                for hcell in head:
                    a('<th>%s</th>' % html.escape(hcell))
                a('</tr></thead><tbody>')
                for row in rows:
                    a('<tr>' + ''.join('<td>%s</td>' % h_inline(c) for c in row) + '</tr>')
                a('</tbody></table></div>')

    # cheat sheet
    a('<h2 id="cheat">%d. Cheat sheet &mdash; memorise this page</h2>' % N_CHE)
    a('<p>Cover the right-hand column and try to write each line from memory.</p>')
    a('<div class="tblwrap"><table class="cheat"><thead><tr><th>You want to&hellip;</th>'
      '<th>Write</th></tr></thead><tbody>')
    for what, how in C.CHEATSHEET:
        a('<tr><td>%s</td><td>%s</td></tr>' % (h_inline(what), h_inline(how)))
    a('</tbody></table></div>')

    # pitfalls
    a('<h2 id="pitfalls">%d. %s</h2>' % (N_PIT, html.escape(PIT_T)))
    a('<div class="tblwrap"><table class="pit"><tbody>')
    for title, expl in C.PITFALLS:
        a('<tr><td><b>%s</b><span class="tech">%s</span></td></tr>'
          % (h_inline(title), h_inline(expl)))
    a('</tbody></table></div>')

    # complexity
    a('<h2 id="complexity">%d. Complexity reference</h2>' % N_CPX)
    a('<div class="tblwrap"><table><thead><tr><th>Operation</th><th>Cost</th></tr></thead><tbody>')
    for op, cost in C.COMPLEXITY:
        a('<tr><td>%s</td><td>%s</td></tr>' % (h_inline(op), h_inline(cost)))
    a('</tbody></table></div>')

    # problems
    a('<h2 id="problems">%d. Practice &mdash; %d problems, Group %s</h2>'
      % (N_PRB, NPROB, PGROUP))
    a('<p>All %d are from the <a href="https://arena.olimpiici.com/#/catalog/1" '
      'style="color:var(--accent)">Olimpiici Arena catalogue</a>, Group %s, across the National '
      'Olympiad rounds and the seasonal tournaments. They are ordered easiest first inside each '
      'group, and the groups themselves get harder as you go down. Tick a problem when it is '
      'accepted &mdash; progress is stored in this browser. The full classified catalogue is in '
      '<a href="../Group C and D problems/index.html" style="color:var(--accent)">'
      'Group C and D problems</a>.</p>' % (NPROB, PGROUP))
    a('<div id="bar"><strong id="cnt"></strong><div class="track"><div class="fill"></div></div>'
      '<button id="reset">Reset</button></div>')

    for key, name, maps in GROUPS:
        a('<div class="pgroup"><h3 id="g%s">%s</h3><p class="maps">%s</p>' % (key, html.escape(name), html.escape(maps)))
        a('<div class="tblwrap"><table class="prob"><thead><tr><th style="width:26px"></th>'
          '<th>Problem</th><th>Technique</th><th>Source</th><th>Level</th></tr></thead><tbody>')
        for p in [x for x in SEL if x['section'] == key]:
            src = '%s %s' % (ROUND_EN.get(p['family'], p['family']), p['year'])
            a('<tr>'
              '<td><input type="checkbox" data-pid="%d"></td>'
              '<td><a href="%s" target="_blank" rel="noopener">%s</a><br>'
              '<span class="tech">%s</span></td>'
              '<td class="tech">%s</td>'
              '<td class="src">%s</td>'
              '<td class="dif">%s</td></tr>'
              % (p['pid'], p['url'], html.escape(p['title']), html.escape(p['en']),
                 h_inline(p['tech']), html.escape(src), stars(p['diff'])))
        a('</tbody></table></div></div>')

    a('<footer><a href="../index.html" style="color:var(--accent)">&larr; All lessons</a>'
      ' &middot; %s</footer>' % html.escape(FOOTER))
    a('</main></div>')
    a('<script>var LESSON_SLUG=%s;var LEGACY_KEY=%s;</script>'
      % (json.dumps(SLUG), json.dumps(getattr(C, "LEGACY_DONE_KEY", None))))
    a('<script>%s</script>' % JS)
    a('</body></html>')

    path = OUT / (SLUG + ".html")
    path.write_text('\n'.join(P), encoding='utf-8')
    return path


# =============================================================== DOCX
def build_docx():
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    ACCENT = RGBColor(0x8A, 0x4B, 0x1F)
    MUTED  = RGBColor(0x60, 0x5C, 0x58)
    CODEINK= RGBColor(0x2D, 0x2A, 0x26)

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12

    for nm, sz, col in (('Heading 1', 17, ACCENT), ('Heading 2', 13, ACCENT), ('Heading 3', 11.5, ACCENT)):
        s = doc.styles[nm]
        s.font.name = 'Calibri'; s.font.size = Pt(sz); s.font.bold = True; s.font.color.rgb = col

    def shade(par, hexcolor):
        el = OxmlElement('w:shd')
        el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), hexcolor)
        par._p.get_or_add_pPr().append(el)

    def left_border(par, hexcolor):
        pbdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:left')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '18')
        b.set(qn('w:space'), '6'); b.set(qn('w:color'), hexcolor)
        pbdr.append(b)
        par._p.get_or_add_pPr().append(pbdr)

    def runs(par, text, size=10.5, base_bold=False):
        for kind, t in inline(text):
            r = par.add_run(t)
            r.font.size = Pt(size)
            r.bold = base_bold
            if kind == 'b':
                r.bold = True
            elif kind == 'c':
                r.font.name = 'Consolas'
                r.font.size = Pt(size - 1)
                r.font.color.rgb = CODEINK
                r._element.rPr.rFonts.set(qn('w:cs'), 'Consolas')
        return par

    def para(text, size=10.5, style=None):
        p = doc.add_paragraph(style=style)
        runs(p, text, size)
        return p

    def code_block(text):
        for i, line in enumerate(text.split('\n')):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(6 if i == 0 else 0)
            pf.space_after  = Pt(6 if i == len(text.split('\n')) - 1 else 0)
            pf.left_indent  = Inches(0.16)
            pf.line_spacing = 1.0
            shade(p, 'F4F1EC')
            left_border(p, '8A4B1F')
            r = p.add_run(line if line else ' ')
            r.font.name = 'Consolas'; r.font.size = Pt(8.5); r.font.color.rgb = CODEINK
            r._element.rPr.rFonts.set(qn('w:cs'), 'Consolas')

    def note_block(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_before = Pt(6)
        shade(p, 'FDF6E7')
        left_border(p, 'D9A441')
        runs(p, text, 10)

    def make_table(head, rows, widths=None, style='Table Grid'):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = style
        t.autofit = True
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ''
            r = c.paragraphs[0].add_run(h.upper())
            r.bold = True; r.font.size = Pt(8); r.font.color.rgb = MUTED
            shade(c.paragraphs[0], 'F3ECE4')
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ''
                runs(cells[i].paragraphs[0], val, 9.5)
        if widths:
            for r_ in t.rows:
                for i, w in enumerate(widths):
                    r_.cells[i].width = Inches(w)
        return t

    def hyperlink(par, url, text):
        r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), r_id)
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        col = OxmlElement('w:color'); col.set(qn('w:val'), '8A4B1F'); rPr.append(col)
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
        f = OxmlElement('w:rFonts'); f.set(qn('w:ascii'), 'Consolas'); f.set(qn('w:hAnsi'), 'Consolas')
        rPr.append(f)
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '19'); rPr.append(sz)
        b = OxmlElement('w:b'); rPr.append(b)
        r.append(rPr)
        t = OxmlElement('w:t'); t.text = text; r.append(t)
        link.append(r); par._p.append(link)

    # ---------------- title
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run(C.TITLE); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = ACCENT
    sp = doc.add_paragraph(); r = sp.add_run(C.SUBTITLE)
    r.font.size = Pt(10.5); r.font.color.rgb = MUTED
    doc.add_paragraph()

    # ---------------- TOC field
    doc.add_heading('Contents', level=2)
    p = doc.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), r'TOC \o "1-2" \h \z \u')
    inner = OxmlElement('w:r'); it = OxmlElement('w:t')
    it.text = 'Right-click here and choose "Update Field" to build the table of contents.'
    inner.append(it); fld.append(inner); p._p.append(fld)
    doc.add_page_break()

    # ---------------- theory
    for sid, title, blocks in C.SECTIONS:
        doc.add_heading(title, level=1)
        for kind, *rest in blocks:
            if kind == 'p':
                para(rest[0])
            elif kind == 'h3':
                doc.add_heading(rest[0], level=3)
            elif kind == 'code':
                code_block(rest[0])
            elif kind == 'note':
                note_block(rest[0])
            elif kind == 'list':
                for item in rest[0]:
                    para(item, style='List Bullet')
            elif kind == 'table':
                head, rows = rest
                w = [2.1, 3.3, 0.9] if len(head) == 3 else [2.3, 4.0]
                make_table(head, rows, w)
                doc.add_paragraph()

    # ---------------- cheat sheet
    doc.add_page_break()
    doc.add_heading('%d. Cheat sheet \u2014 memorise this page' % N_CHE, level=1)
    para('Cover the right-hand column and try to write each line from memory.')
    make_table(['You want to\u2026', 'Write'], [[a_, b_] for a_, b_ in C.CHEATSHEET], [2.2, 4.1])

    # ---------------- pitfalls
    doc.add_page_break()
    doc.add_heading('%d. %s' % (N_PIT, PIT_T), level=1)
    for i, (title, expl) in enumerate(C.PITFALLS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run('%d. ' % i); r.bold = True; r.font.color.rgb = ACCENT
        runs(p, title, 10.5, base_bold=True)
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Inches(0.22)
        q.paragraph_format.space_after = Pt(8)
        runs(q, expl, 10)

    # ---------------- complexity
    doc.add_heading('%d. Complexity reference' % N_CPX, level=1)
    make_table(['Operation', 'Cost'], C.COMPLEXITY, [4.2, 2.1])

    # ---------------- problems
    doc.add_page_break()
    doc.add_heading('%d. Practice \u2014 %d problems, Group %s' % (N_PRB, NPROB, PGROUP), level=1)
    para('All %d problems are from the Olimpiici Arena catalogue '
         '(https://arena.olimpiici.com/#/catalog/1), Group %s, across the National Olympiad rounds '
         'and the seasonal tournaments. They are ordered easiest first inside each group, and the '
         'groups themselves get harder as you go down. Click a problem name to open it.'
         % (NPROB, PGROUP))

    for key, name, maps in GROUPS:
        doc.add_heading(name, level=2)
        p = doc.add_paragraph(); r = p.add_run(maps)
        r.font.size = Pt(9); r.font.color.rgb = MUTED; r.italic = True
        p.paragraph_format.space_after = Pt(4)

        t = doc.add_table(rows=1, cols=4)
        t.style = 'Table Grid'
        for i, hcell in enumerate(['Problem', 'Technique', 'Source', 'Level']):
            c = t.rows[0].cells[i]; c.text = ''
            rr = c.paragraphs[0].add_run(hcell.upper())
            rr.bold = True; rr.font.size = Pt(8); rr.font.color.rgb = MUTED
            shade(c.paragraphs[0], 'F3ECE4')
        for p_ in [x for x in SEL if x['section'] == key]:
            cells = t.add_row().cells
            cells[0].text = ''
            hyperlink(cells[0].paragraphs[0], p_['url'], p_['title'])
            sub = cells[0].add_paragraph()
            sub.paragraph_format.space_before = Pt(0); sub.paragraph_format.space_after = Pt(0)
            rr = sub.add_run(p_['en']); rr.font.size = Pt(8.5); rr.font.color.rgb = MUTED
            cells[1].text = ''
            runs(cells[1].paragraphs[0], p_['tech'], 9)
            cells[2].text = ''
            rr = cells[2].paragraphs[0].add_run(
                '%s\n%s' % (ROUND_EN.get(p_['family'], p_['family']), p_['year']))
            rr.font.size = Pt(8.5); rr.font.color.rgb = MUTED
            cells[3].text = ''
            rr = cells[3].paragraphs[0].add_run(stars(p_['diff']))
            rr.font.size = Pt(9); rr.font.color.rgb = ACCENT
        for r_ in t.rows:
            for i, w in enumerate([1.55, 2.5, 1.35, 0.85]):
                r_.cells[i].width = Inches(w)
        doc.add_paragraph()

    path = OUT / (SLUG + ".docx")
    doc.save(str(path))
    return path


if __name__ == '__main__':
    hp = build_html()
    print('HTML ->', hp, hp.stat().st_size, 'bytes')
    if not _args.html_only:
        dp = build_docx()
        print('DOCX ->', dp, dp.stat().st_size, 'bytes')
